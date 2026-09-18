# -*- coding: utf-8 -*-
"""合同生成引擎：复制模板并替换占位符，输出 Word。"""
import copy
import os
import re
from datetime import datetime

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

PLACEHOLDER_RE = re.compile(r"\{([^}]+)\}")
SIGNATURE_TAB_POS_TWIPS = 4828

TEMPLATE_MAP = {
    "房屋租赁合同": "房屋租赁合同样板.docx",
    "管理服务合同": "管理服务合同模板.docx",
}


def fmt_date(value):
    if not value:
        return ""
    s = str(value)[:10]
    try:
        d = datetime.strptime(s, "%Y-%m-%d")
    except ValueError:
        return s
    return "%d年%d月%d日" % (d.year, d.month, d.day)


def fmt_date_from_ts(v):
    if v is None or v == "":
        return ""
    if isinstance(v, (int, float)):
        try:
            d = datetime.fromtimestamp(v / 1000)
            return "%d年%d月%d日" % (d.year, d.month, d.day)
        except Exception:
            return str(v)
    return fmt_date(v)


def to_text(v):
    if v is None:
        return ""
    if isinstance(v, list):
        parts = []
        for item in v:
            if isinstance(item, dict):
                parts.append(str(item.get("text", item.get("name", ""))))
            else:
                parts.append(str(item))
        return "、".join(parts)
    if isinstance(v, float) and v == int(v):
        return str(int(v))
    if isinstance(v, bool):
        return "是" if v else "否"
    if isinstance(v, dict):
        return str(v.get("text", v.get("name", "")))
    return str(v)


def replace_in_paragraph(p, mapping):
    runs = p.runs
    if not runs:
        return False
    full = "".join(r.text for r in runs)
    if "{" not in full:
        return False
    new_full = full
    found = False
    for ph in PLACEHOLDER_RE.findall(full):
        if ph in mapping:
            new_full = new_full.replace("{" + ph + "}", str(mapping[ph]))
            found = True
    if not found:
        return False
    first = runs[0]
    for r in runs[1:]:
        r.text = ""
    lines = str(new_full).split("\n")
    first.text = lines[0]
    for line in lines[1:]:
        first.add_break()
        first.add_text(line)
    return True


def check_leftover(doc):
    leftovers = set()
    texts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                texts.extend(p.text for p in cell.paragraphs)
    for t in texts:
        for ph in PLACEHOLDER_RE.findall(t):
            leftovers.add(ph)
    return leftovers


def _rpr_without_underline(run):
    rpr = run._r.find(qn("w:rPr"))
    if rpr is None:
        return None
    new_rpr = copy.deepcopy(rpr)
    u = new_rpr.find(qn("w:u"))
    if u is not None:
        new_rpr.remove(u)
    return new_rpr


def _add_left_tab_stop(paragraph, pos_twips):
    pPr = paragraph._p.get_or_add_pPr()
    tabs = pPr.find(qn("w:tabs"))
    if tabs is None:
        tabs = OxmlElement("w:tabs")
        pPr.append(tabs)
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "left")
    tab.set(qn("w:pos"), str(int(pos_twips)))
    tabs.append(tab)


def fix_signature_block(doc, mapping):
    d = mapping.get("租赁起始日期") or ""
    for p in doc.paragraphs:
        stripped = p.text.replace(" ", "").replace("\u3000", "")
        if stripped == "甲方：乙方：":
            parts = ["甲方：", "乙方："]
        elif stripped == "甲方代理人：乙方代理人：":
            parts = ["甲方代理人：", "乙方代理人："]
        elif d and stripped == d + d:
            parts = [d, d]
        else:
            continue
        rpr = _rpr_without_underline(p.runs[0]) if p.runs else None
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        _add_left_tab_stop(p, SIGNATURE_TAB_POS_TWIPS)
        run1 = p.add_run(parts[0])
        run_tab = p.add_run()
        run_tab._r.append(OxmlElement("w:tab"))
        run2 = p.add_run(parts[1])
        for r in (run1, run_tab, run2):
            if rpr is not None:
                r._r.insert(0, copy.deepcopy(rpr))


def indent_account_info(doc):
    for p in doc.paragraphs:
        if p.text.startswith("账户名称："):
            pPr = p._p.get_or_add_pPr()
            ind = pPr.find(qn("w:ind"))
            if ind is None:
                ind = OxmlElement("w:ind")
                pPr.append(ind)
            ind.set(qn("w:left"), "480")
            ind.set(qn("w:leftChars"), "200")
            ind.set(qn("w:firstLine"), "0")
            ind.set(qn("w:firstLineChars"), "0")
            return True
    return False


def generate_one(rec, template_dir, output_path):
    ctype = to_text(rec.get("合同类型"))
    tmpl = TEMPLATE_MAP.get(ctype)
    if not tmpl:
        return {"ok": False, "原因": "未知合同类型: %s" % ctype}
    src = os.path.join(template_dir, tmpl)
    if not os.path.exists(src):
        return {"ok": False, "原因": "模板不存在: %s" % src}
    mapping = {
        "企业名称": to_text(rec.get("企业名称")),
        "收款抬头": to_text(rec.get("收款抬头")),
        "地址": to_text(rec.get("地址")),
        "房间号": to_text(rec.get("房间号")),
        "使用面积": to_text(rec.get("使用面积")),
        "地址费": to_text(rec.get("地址费")),
        "租赁起始日期": fmt_date_from_ts(rec.get("租赁起始日期")),
        "租赁截止日期": fmt_date_from_ts(rec.get("租赁截止日期")),
        "大写金额": to_text(rec.get("大写金额")),
        "收款信息": to_text(rec.get("收款信息")),
    }
    doc = Document(src)
    count = 0
    for p in doc.paragraphs:
        if replace_in_paragraph(p, mapping):
            count += 1
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if replace_in_paragraph(p, mapping):
                        count += 1
    leftovers = check_leftover(doc)
    fix_signature_block(doc, mapping)
    if ctype == "管理服务合同":
        indent_account_info(doc)
    doc.save(output_path)
    return {
        "ok": True,
        "企业名称": mapping["企业名称"],
        "合同类型": ctype,
        "替换处数": count,
        "未替换占位符": sorted(leftovers),
        "mapping": mapping,
    }
